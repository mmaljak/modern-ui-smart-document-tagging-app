"""Run once to generate invoice_sample.pdf and contract_sample.docx.

Usage:
    pip install fpdf2 python-docx
    python docs/samples/create_binary_samples.py
"""
from __future__ import annotations

import pathlib

SAMPLES_DIR = pathlib.Path(__file__).parent


def create_invoice_pdf() -> None:
    from fpdf import FPDF  # type: ignore

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "INVOICE", ln=True, align="C")
    pdf.ln(5)
    pdf.set_font("Helvetica", size=11)
    lines = [
        "Invoice Number: INV-2024-001",
        "Invoice Date: January 15, 2024",
        "Due Date: February 14, 2024",
        "",
        "Bill To: Acme Corporation, 123 Business Ave, New York, NY 10001",
        "From: TechConsult LLC, 456 Maple Street, San Francisco, CA 94102",
        "",
        "DESCRIPTION OF SERVICES",
        "Software Architecture Consulting  40hrs @ $150.00/hr  $6,000.00",
        "Code Review and QA               10hrs @ $120.00/hr  $1,200.00",
        "Technical Documentation           8hrs @ $100.00/hr  $  800.00",
        "",
        "Subtotal:          $8,000.00",
        "Tax (8%):          $  640.00",
        "Total Amount Due:  $8,640.00",
        "",
        "Payment Terms: Net 30",
        "Please remit payment by the due date.",
        "Late payments are subject to a 1.5% monthly finance charge.",
    ]
    for line in lines:
        pdf.cell(0, 7, line, ln=True)

    out = SAMPLES_DIR / "invoice_sample.pdf"
    pdf.output(str(out))
    print(f"Created {out}")


def create_contract_docx() -> None:
    from docx import Document  # type: ignore
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("SERVICE AGREEMENT", 0)

    paragraphs = [
        "This Service Agreement ('Agreement') is entered into as of March 1, 2024, "
        "by and between TechConsult LLC ('Service Provider') and Acme Corporation ('Client').",
        "",
        "WHEREAS, Service Provider desires to provide technology services to Client; and",
        "WHEREAS, Client desires to retain Service Provider for such services;",
        "",
        "1. SERVICES",
        "Service Provider shall perform software development and consulting services as agreed.",
        "",
        "2. OBLIGATIONS",
        "Service Provider shall perform Services professionally. Client shall pay invoices Net 30.",
        "",
        "3. TERMINATION",
        "Either party may terminate this Agreement upon 30 days written notice. "
        "In the event of material breach, the non-breaching party may terminate immediately.",
        "",
        "4. GOVERNING LAW",
        "This Agreement shall be governed by the laws of the State of Delaware.",
        "",
        "IN WITNESS WHEREOF, the parties have executed this Agreement.",
        "",
        "TECHCONSULT LLC                         ACME CORPORATION",
        "By: ___________________                 By: ___________________",
        "Name: Jane Smith                        Name: John Doe",
        "Title: Managing Director                Title: Chief Technology Officer",
    ]

    for para in paragraphs:
        p = doc.add_paragraph(para)
        p.runs[0].font.size = Pt(11) if para else None

    out = SAMPLES_DIR / "contract_sample.docx"
    doc.save(str(out))
    print(f"Created {out}")


if __name__ == "__main__":
    create_invoice_pdf()
    create_contract_docx()
    print("Done.")
